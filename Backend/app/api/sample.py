"""A08 샘플 BOM CSV.

프론트 mock(Frontend/src/api/mock/data/session.mock.json)의 rows[].raw 4필드를 그대로 옮긴 것이다(서버가 프론트 파일을 읽지 않도록 복사).
가상 데이터이며 자동차 부품 BOM 으로 교체할 예정이다(교체는 이 작업 범위가 아님).
"""
import csv
import io

from fastapi import APIRouter
from fastapi.responses import Response

router = APIRouter(tags=["sample"])

SAMPLE_ROWS = [
    ["SGCC 용융아연도금강판", "0.8T", "kg", "1.1"],
    ["ABS 수지 펠릿", "", "kg", "0.12"],
    ["재생 알루미늄 압출재 A6063", "", "kg", "0.9"],
    ["PC 수지", "", "kg", "0.07"],
    ["점용접 브라켓 어셈블리", "", "ea", "4"],
    ["분체도장 처리", "", "회", "1"],
    ["SPCC 냉연강판 1.2T", "t1.2 x 120 x 340", "kg", "2.4"],
    ["SPCC 냉연강판 1.2T", "t1.2 x 80 x 200", "kg", "1.6"],
    ["고밀도 폴리에틸렌 HDPE", "", "kg", "0.35"],
    ["볼트 M6x20", "", "EA", "20"],
    ["PP 사출품 커버", "", "kg", "0.22"],
    ["알루미늄 판재 A5052", "", "m2", "0.8"],
]


@router.get("/sample.csv")
def sample_csv():  # A08
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(["품명", "규격", "단위", "수량"])
    w.writerows(SAMPLE_ROWS)
    # 맨 앞 BOM: 엑셀에서 한글이 깨지지 않게
    return Response(("\ufeff" + buf.getvalue()).encode("utf-8"), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="sample_bom.csv"'})


TEMPLATE_HEADER = ["부품번호", "구분", "품명", "규격", "단위", "수량"]  # 앞의 두 열은 선택(없어도 된다)
TEMPLATE_EXAMPLES = [  # 예시 3행(회색 글씨, "예시" 메모). 지우고 쓰도록 안내한다
    ["P01", "소재", "SPCC 냉연강판", "t1.2 x 120 x 340", "kg", "2.4"],
    ["P01", "공정", "프레스 가공", "", "kg", "2.4"],
    ["P02", "소재", "알루미늄 판재 A5052", "t2.0", "m2", "0.8"],
]
FIELD_GUIDE = [
    ["부품번호", "선택", "같은 부품의 행을 알아보기 위한 번호. 지금은 결과 파일에 그대로 적을 뿐 검색에는 쓰지 않습니다."],
    ["구분", "선택", "소재 또는 공정(재료·원자재·자재·material → 소재, 가공·process → 공정). 비우면 소재·공정 모두에서 찾습니다."],
    ["품명", "필수", "자재·부품 이름. 약어·영문도 됩니다(예: SPCC 냉연강판)."],
    ["규격", "선택", "두께·치수·재질 등. 검색 질의에 품명과 함께 들어갑니다."],
    ["단위", "필수", "아래 허용 단위 중 하나. 그 밖의 단위는 '단위 확인 필요'로 표시됩니다."],
    ["수량", "필수", "숫자. 천 단위 쉼표(1,200)는 읽습니다. 0 이하는 '수량 이상'으로 표시됩니다."],
]


@router.get("/template.xlsx")
def template_xlsx():  # A08b
    import openpyxl
    from openpyxl.comments import Comment
    from openpyxl.styles import Font

    from app.bom.clean import UNITS
    from app.bom.reader import MAX_ROWS

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "BOM"  # 리더가 이 이름의 시트를 먼저 읽는다
    ws.append(TEMPLATE_HEADER)
    for col in (1, 2):  # 머리글 이름은 그대로 두어야 올렸을 때 자동으로 인식된다. 선택 표시는 메모로 한다
        ws.cell(row=1, column=col).comment = Comment("선택 — 비워도 됩니다", "LCI-Match")
    for row in TEMPLATE_EXAMPLES:
        ws.append(row)
        for cell in ws[ws.max_row]:
            cell.font = Font(color="808080")
        ws.cell(row=ws.max_row, column=3).comment = Comment("예시 — 지우고 실제 BOM 을 적으세요", "LCI-Match")
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEF", (10, 8, 28, 22, 8, 10)):
        ws.column_dimensions[col].width = width

    guide = wb.create_sheet("작성 안내")
    guide.append(["필드", "필수 여부", "설명"])
    for row in FIELD_GUIDE:
        guide.append(row)
    guide.append([])
    guide.append(["허용 단위(정규)", "화면 표기", "입력할 수 있는 표기(대소문자·공백 무시)"])
    by_unit = {}
    for key, (unit, display) in UNITS.items():  # clean 단위 표를 그대로 쓴다
        by_unit.setdefault((unit, display), []).append(key)
    for (unit, display), keys in by_unit.items():
        guide.append([unit, display or "입력한 그대로", ", ".join(keys)])
    guide.append([])
    guide.append(["상한", "", f"파일 5MB, 데이터 {MAX_ROWS:,}행까지"])
    guide.append(["작성 규칙", "", "한 행에 자재 하나. 첫 줄은 머리행(부품번호·구분·품명·규격·단위·수량). 시트 이름 BOM 을 먼저 읽습니다."])
    for col, width in zip("ABC", (18, 12, 70)):
        guide.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=\"BOM_template.xlsx\"; filename*=UTF-8''BOM_%EC%96%91%EC%8B%9D.xlsx"})
