"""BOM 파일 읽기(.csv, .xlsx, .xls). 반환: (헤더, [(source_row, 값 목록)]). source_row 는 원본 파일의 줄/행 번호(헤더 포함 1부터)."""
import csv
import io
from datetime import date, datetime
from pathlib import PurePath

from app.bom.mapping import ALIASES, norm
from app.errors import ApiError

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
DELIMITERS = (",", ";", "\t")
SHEET_FIRST = ("BOM", "자재명세서")  # 이 이름의 시트를 먼저 읽는다
SHEET_SKIP = ("작성방법", "예제")  # 안내용 시트: 첫 시트로 고르지 않는다
HEADER_SCAN = 30  # 안내 문구 아래의 머리행을 찾을 때 훑는 줄 수
HEADER_MIN_HITS = 3  # 별칭 사전에 맞는 칸이 이만큼 있으면 머리행
ALL_ALIASES = set().union(*ALIASES.values())
# FTA 양식의 기본 예시 2행(품번, 품명(영문), 품명(국문), 규격, 단위, 소요량). 이것과 완전히 같을 때만 예시로 보고 뺀다
EXAMPLE_ROWS = [("AAA-01", "COVER", "커버", "cover", "EA", 1.0), ("AAA-02", "CASE", "케이스", "450X450X200mm", "EA", 1.0)]
EXAMPLE_COLS = ("품번", "품명(영문)", "품명(국문)", "규격", "단위", "소요량")


def unreadable(why=""):
    return ApiError(400, "파일을 읽을 수 없습니다. CSV, XLSX 또는 XLS 파일인지 확인해 주세요." + (f" ({why})" if why else ""), "UNREADABLE_FILE")


def count_outside_quotes(line, ch):
    n, quoted = 0, False
    for c in line:
        if c == '"':
            quoted = not quoted
        elif c == ch and not quoted:
            n += 1
    return n


def decode(data):
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            return data.decode("cp949")
        except UnicodeDecodeError as e:
            raise unreadable("인코딩을 알 수 없음") from e


def read_csv(data):
    text = decode(data)
    first = next((line for line in text.splitlines() if line.strip()), "")
    # 첫 줄에서 가장 많이 나온 구분자(같으면 쉼표 > 세미콜론 > 탭). 따옴표 안은 세지 않는다
    delimiter = max(DELIMITERS, key=lambda d: count_outside_quotes(first, d))
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    out = []
    try:
        for record in reader:
            out.append((reader.line_num, record))  # 여러 줄에 걸친 따옴표 칸이면 그 레코드의 마지막 줄 번호
    except csv.Error as e:
        raise unreadable(str(e)) from e
    return out


def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return str(v)


def pick_sheet(names):
    """읽을 시트 이름: 이름이 맞는 시트, 없으면 안내용이 아닌 첫 시트, 그것도 없으면 첫 시트."""
    return next((n for n in SHEET_FIRST if n in names), None) or next((n for n in names if n not in SHEET_SKIP), names[0])


def read_xls(data):
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=data)
        sheet = book.sheet_by_name(pick_sheet(book.sheet_names()))
    except Exception as e:  # xlrd 도 깨진 파일에 여러 종류 예외를 낸다
        raise unreadable("XLS 를 열 수 없음") from e
    out = []
    for i in range(sheet.nrows):
        row = []
        for c in sheet.row(i):
            if c.ctype == xlrd.XL_CELL_DATE:
                try:
                    row.append(xlrd.xldate.xldate_as_datetime(c.value, book.datemode).isoformat())
                except Exception:  # 날짜로 못 바꾸면 숫자 그대로
                    row.append(cell_text(c.value))
            elif c.ctype in (xlrd.XL_CELL_ERROR, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_EMPTY):
                row.append("")
            else:
                row.append(cell_text(c.value))
        out.append((i + 1, row))
    return out


def read_xlsx(data):
    import openpyxl
    try:
        book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)  # 수식은 저장된 값으로
    except Exception as e:  # openpyxl 은 깨진 파일에 여러 종류 예외를 낸다
        raise unreadable("XLSX 를 열 수 없음") from e
    try:
        ws = book[pick_sheet(book.sheetnames)]
        # read_only 모드에서 병합 셀은 왼쪽 위에만 값이 있고 나머지는 None 이다
        return [(i, [cell_text(v) for v in row]) for i, row in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1)]
    finally:
        book.close()


def unique_header(cells):
    """앞뒤 공백 제거. 같은 이름이 또 나오면 뒤에 .2, .3 … 을 붙인다."""
    seen, out = {}, []
    for c in cells:
        name = c.strip()
        seen[name] = seen.get(name, 0) + 1
        out.append(name if seen[name] == 1 else f"{name}.{seen[name]}")
    return out


def find_header(records):
    """머리행의 위치. 앞 HEADER_SCAN 줄 중 별칭에 맞는 칸이 HEADER_MIN_HITS 개 이상인 첫 줄, 없으면 0(첫 줄)."""
    for k, (_, r) in enumerate(records[:HEADER_SCAN]):
        if sum(1 for c in r if norm(c) in ALL_ALIASES) >= HEADER_MIN_HITS:
            return k
    return 0


def name_columns(header):
    """품명 별칭에 맞는 열의 위치 목록."""
    return [i for i, h in enumerate(header) if norm(h) in ALIASES["품명"]]


def is_example(header, row):
    """양식 기본 예시 2행과 모든 칸이 같은가."""
    if any(c not in header for c in EXAMPLE_COLS):
        return False
    vals = [row[header.index(c)].strip() for c in EXAMPLE_COLS]
    try:
        qty = float(vals[5])
    except ValueError:
        return False
    return any(tuple(vals[:5]) == e[:5] and qty == e[5] for e in EXAMPLE_ROWS)


def read_bom_ex(filename, data):
    """(헤더, 행, 뺀 행 수 {"example": n, "empty_name": m}). 뺀 행은 FTA 양식(안내 문구 아래 머리행, 또는 품명 영문·국문 두 열)일 때만 센다."""
    if len(data) > MAX_BYTES:
        raise ApiError(413, "파일은 5MB 이하여야 합니다.", "FILE_TOO_LARGE")
    ext = PurePath(filename or "").suffix.lower()
    if ext == ".csv":
        records = read_csv(data)
    elif ext == ".xlsx":
        records = read_xlsx(data)
    elif ext == ".xls":
        records = read_xls(data)
    else:
        raise unreadable(f"지원하지 않는 확장자 {ext or '(없음)'}")
    records = [(n, r) for n, r in records if any(c.strip() for c in r)]  # 빈 줄은 버린다
    if not records:
        raise unreadable("빈 파일")
    start = find_header(records)
    _, head = records[start]
    while head and not head[-1].strip():  # 헤더 오른쪽 빈 칸은 열로 치지 않는다
        head = head[:-1]
    header = unique_header(head)
    rows = [(n, [(r[i] if i < len(r) else "") for i in range(len(header))]) for n, r in records[start + 1:]]
    skipped = {"example": 0, "empty_name": 0}
    ko, en = (header.index(c) if c in header else None for c in ("품명(국문)", "품명(영문)"))
    fta = start > 0 or (ko is not None and en is not None)
    if fta:
        kept = []
        for n, r in rows:
            if is_example(header, r):
                skipped["example"] += 1
                continue
            if ko is not None and en is not None and not r[ko].strip():  # 국문이 비면 영문으로 대신한다
                r[ko] = r[en]
            if not any(r[i].strip() for i in name_columns(header)):
                skipped["empty_name"] += 1
                continue
            kept.append((n, r))
        rows = kept
    if not header or not rows:
        raise unreadable("헤더 또는 데이터 행이 없음")
    if len(rows) > MAX_ROWS:
        raise ApiError(400, f"행이 너무 많습니다({len(rows):,}행). 한 번에 {MAX_ROWS:,}행까지 올릴 수 있습니다.", "TOO_MANY_ROWS",
                       {"rows": len(rows), "limit": MAX_ROWS})
    return header, rows, skipped


def read_bom(filename, data):
    header, rows, _ = read_bom_ex(filename, data)
    return header, rows
