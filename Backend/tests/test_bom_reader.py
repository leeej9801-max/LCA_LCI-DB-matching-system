"""BOM 파일 읽기 시험(DB 없음)."""
import io

import openpyxl
import pytest

from app.bom.reader import MAX_BYTES, MAX_ROWS, read_bom, read_bom_ex
from app.errors import ApiError


def err(fn):
    with pytest.raises(ApiError) as e:
        fn()
    return e.value


def test_csv_utf8_bom_and_cp949_and_delimiters():
    text = "품명,단위,수량\n냉연강판,kg,2.4\n"
    for data in (text.encode("utf-8"), ("\ufeff" + text).encode("utf-8"), text.encode("cp949")):
        assert read_bom("a.csv", data) == (["품명", "단위", "수량"], [(2, ["냉연강판", "kg", "2.4"])])
    for d in (";", "\t"):
        assert read_bom("a.csv", text.replace(",", d).encode("utf-8"))[1] == [(2, ["냉연강판", "kg", "2.4"])]


def test_delimiter_ignores_quoted_commas_and_source_row_counts_blank_lines():
    data = 'name;spec;"a,b,c,d"\n\nX;"1,2";3\n'.encode("utf-8")  # 따옴표 안 쉼표 4개는 세지 않는다 → 세미콜론
    header, rows = read_bom("a.csv", data)
    assert header == ["name", "spec", "a,b,c,d"] and rows == [(3, ["X", "1,2", "3"])]  # 빈 2행을 건너뛰어도 원본 줄 번호 3


def test_duplicate_header_and_short_rows():
    header, rows = read_bom("a.csv", " 품명 ,품명,단위,품명\nA,B\n".encode("utf-8"))
    assert header == ["품명", "품명.2", "단위", "품명.3"] and rows == [(2, ["A", "B", "", ""])]


def xlsx(sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_bom_sheet_or_first_sheet():
    data = xlsx([("표지", [["제목"]]), ("BOM", [[None], ["품명", "수량"], ["강판", 2.0], ["볼트", 1.5]])])
    assert read_bom("a.xlsx", data) == (["품명", "수량"], [(3, ["강판", "2"]), (4, ["볼트", "1.5"])])  # 1행 빈 줄, 숫자 2.0 → "2"
    data = xlsx([("Sheet1", [["품명"], ["A"]]), ("기타", [["x"], ["y"]])])
    assert read_bom("a.XLSX", data) == (["품명"], [(2, ["A"])])


def test_errors():
    assert err(lambda: read_bom("a.csv", b"x" * (MAX_BYTES + 1))).code == "FILE_TOO_LARGE"
    assert err(lambda: read_bom("a.xls", b"abc")).code == "UNREADABLE_FILE"
    assert err(lambda: read_bom("a.csv", b"")).code == "UNREADABLE_FILE"
    assert err(lambda: read_bom("a.csv", "품명\n\n".encode())).code == "UNREADABLE_FILE"  # 헤더만
    assert err(lambda: read_bom("a.xlsx", b"not a zip")).code == "UNREADABLE_FILE"
    big = ("품명\n" + "a\n" * (MAX_ROWS + 1)).encode()
    e = err(lambda: read_bom("a.csv", big))
    assert (e.status, e.code, e.meta) == (400, "TOO_MANY_ROWS", {"rows": MAX_ROWS + 1, "limit": MAX_ROWS})
    assert len(read_bom("a.csv", ("품명\n" + "a\n" * MAX_ROWS).encode())[1]) == MAX_ROWS  # 2,000행은 통과


# --- U6: FTA 원산지 BOM 양식 ---
from pathlib import Path

from app.api.sample import TEMPLATE_DIR, bom_template
from app.bom.mapping import detect

FTA_HEAD = ["품번", "품명(영문)", "품명(국문)", "규격", "세번(HS CODE)", "단가", "단위", "소요량", "원산지"]
MEASURE = Path(__file__).resolve().parents[2] / "docs" / "LCA" / "data" / "fta" / "measure-v1-input.xlsx"


def fta_rows(*body):
    notes = [[None, "소요부품(자재) 명세서"], [], [None, "□ 완제품 품명(영문/국문) :"], [None, "□ 원재료 사용내역"], []]
    return notes + [[None] + FTA_HEAD] + [[None] + list(b) for b in body]


def fields(header, row):
    m = {x["field"]: x["source"] for x in detect(header, row, {})}
    return {f: row[header.index(s)] for f, s in m.items() if s}


def test_fta_template_excludes_two_examples_and_reports_count():
    # 예시 2행만 있으니 읽을 행이 없어 에러가 되므로, 건수는 예시 아래에 실제 행을 하나 붙여서 확인한다
    wb = openpyxl.load_workbook(TEMPLATE_DIR / "fta-bom-template.xlsx")
    wb["자재명세서"].append([None, "P1", None, "강판", None, None, None, "kg", 1])
    buf = io.BytesIO()
    wb.save(buf)
    header, rows, skipped = read_bom_ex("t.xlsx", buf.getvalue())
    assert skipped == {"example": 2, "empty_name": 0} and len(rows) == 1


def test_fta_example_only_file_has_no_data_rows():
    # 예시 2행만 있는 양식은 읽을 행이 없다(에러). 예시 건수는 위 시험에서 확인한다
    assert err(lambda: read_bom("t.xlsx", (TEMPLATE_DIR / "fta-bom-template.xlsx").read_bytes())).code == "UNREADABLE_FILE"


def test_fta_example_must_match_exactly():
    e = ["AAA-01", "COVER", "커버", "cover", "0000.00", "2000", "EA", "1", "KR"]
    changed = ["AAA-02", "CASE", "케이스", "450X450X200mm", "0000.00", "4000", "EA", "5", "KR"]  # 소요량이 다르다
    real = ["P01", None, "강판", "1T", None, None, "kg", 2, None]
    header, rows, skipped = read_bom_ex("a.xlsx", xlsx([("자재명세서", fta_rows(e, changed, real))]))
    assert skipped == {"example": 1, "empty_name": 0}
    assert [r[header.index("품번")] for _, r in rows] == ["AAA-02", "P01"]


def test_fta_sample_19_rows():
    header, rows, skipped = read_bom_ex("s.xlsx", (TEMPLATE_DIR / "fta-bom-sample.xlsx").read_bytes())
    assert len(rows) == 19 and skipped == {"example": 0, "empty_name": 0}
    f = fields(header, rows[0][1])
    assert f == {"품명": "SGCC 용융아연도금강판", "규격": "0.8T", "단위": "kg", "수량": "1.1", "부품번호": "P01"}
    assert rows[0][0] == 9  # 원본 파일의 행 번호(머리행 8)
    assert all(fields(header, r)["품명"] and fields(header, r)["단위"] and fields(header, r)["수량"] for _, r in rows)


def test_fta_measure_input_115_rows_english_fallback():
    if not MEASURE.exists():
        pytest.skip("docs/ 는 gitignore 라 저장소에 없을 수 있다")
    header, rows, skipped = read_bom_ex("m.xlsx", MEASURE.read_bytes())
    assert len(rows) == 115 and skipped == {"example": 0, "empty_name": 0}
    names = [fields(header, r)["품명"] for _, r in rows]
    assert all(names) and "Natural rubber latex" in names and "열연강판 SPHC" in names


def test_fta_name_columns_prefer_korean_and_fall_back_per_row():
    body = [["P1", "STEEL", "강판", "1T", None, None, "kg", 1, None], ["P2", "Rubber", None, None, None, None, "kg", 2, None]]
    header, rows = read_bom("a.xlsx", xlsx([("자재명세서", fta_rows(*body))]))
    assert [fields(header, r)["품명"] for _, r in rows] == ["강판", "Rubber"]


def test_fta_empty_name_and_tail_rows_are_skipped_and_counted():
    body = [["P1", None, "강판", None, None, None, "kg", 1, None], ["역내산", None, None, None, None, None, None, None, "3"]]
    header, rows, skipped = read_bom_ex("a.xlsx", xlsx([("자재명세서", fta_rows(*body))]))
    assert len(rows) == 1 and skipped == {"example": 0, "empty_name": 1}


def test_header_found_below_notes_for_csv_and_plain_files_unchanged():
    text = "소요부품 명세서\n\n품번,품명,단위,수량\nP1,강판,kg,2\n"
    header, rows = read_bom("a.csv", text.encode("utf-8"))
    assert header == ["품번", "품명", "단위", "수량"] and rows == [(4, ["P1", "강판", "kg", "2"])]
    # 머리행이 첫 줄인 일반 파일, 품명이 빈 행은 그대로 남는다(제외는 파이프라인이 한다)
    header, rows, skipped = read_bom_ex("a.csv", "품명,단위,수량\n,kg,1\n강판,kg,2\n".encode())
    assert len(rows) == 2 and skipped == {"example": 0, "empty_name": 0}
    # 별칭 칸이 3개 미만이면 첫 줄이 머리행
    assert read_bom("a.csv", "품명,수량\nA,1\n".encode())[0] == ["품명", "수량"]


def test_no_role_column_is_not_an_error():
    header, rows = read_bom("a.xlsx", xlsx([("자재명세서", fta_rows(["P1", None, "강판", None, None, None, "kg", 1, None]))]))
    m = {x["field"]: x["method"] for x in detect(header, rows[0][1], {})}
    assert m["구분"] == "missing" and m["품명"] == "auto"


def test_sheet_choice_skips_guide_sheets():
    data = xlsx([("작성방법", [["안내"]]), ("자재명세서", [["품명", "수량"], ["A", 1]])])
    assert read_bom("a.xlsx", data)[0] == ["품명", "수량"]
    data = xlsx([("작성방법", [["안내"]]), ("Sheet1", [["품명", "수량"], ["A", 1]])])
    assert read_bom("a.xlsx", data)[0] == ["품명", "수량"]


def test_xls_read():
    xlwt = pytest.importorskip("xlwt")  # xlwt 는 의존성이 아니다. 없으면 건너뛴다(실파일은 수동 확인)
    wb = xlwt.Workbook()
    ws = wb.add_sheet("자재명세서")
    for i, r in enumerate(fta_rows(["P1", "STEEL", "강판", "1T", None, None, "kg", 2.5, None])):
        for j, v in enumerate(r):
            if v is not None:
                ws.write(i, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    header, rows = read_bom("a.xls", buf.getvalue())
    assert fields(header, rows[0][1]) == {"품명": "강판", "규격": "1T", "단위": "kg", "수량": "2.5", "부품번호": "P1"}


def test_template_endpoint_returns_files():
    for kind, name in (("blank", "fta-bom-template.xlsx"), ("sample", "fta-bom-sample.xlsx")):
        r = bom_template(kind)
        assert r.body == (TEMPLATE_DIR / name).read_bytes() and name in r.headers["content-disposition"]
        assert openpyxl.load_workbook(io.BytesIO(r.body)).sheetnames == ["자재명세서", "작성방법"]
