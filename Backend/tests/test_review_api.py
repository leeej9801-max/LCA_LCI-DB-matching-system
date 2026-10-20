"""검토 흐름 시험: A03 행 변경, A04 설명, A07 내보내기, A08b 양식, review_decision 추가 전용, 0005 기본값. 임시 DB + 가짜 임베더."""
import csv
import io
import shutil
import threading
import time

import openpyxl
import psycopg
import pytest
from psycopg import errors

from app import jobs
from app.bom.reader import read_bom
from app.db.migrate import MIGRATIONS, migrate
from app.schemas import Explanation, RowUpdateResponse, SessionPayload
from test_bom_match import bomdb  # noqa: F401  (fixture 재사용)
from test_indexer import FakeEmbedder
from test_migrate import needs_db, tempdb  # noqa: F401
from test_sessions_api import assert_error, client, post  # noqa: F401

CSV = "품명,단위,수량\n냉연강판,kg,2.4\n볼트,EA,20\n알루미늄 판,m2,0.8\n모르는 것,회,1\n,kg,1\n".encode("utf-8")


@pytest.fixture
def sess(client):  # noqa: F811
    """A01 로 세션을 만들고 작업을 끝낸 뒤 (client, session_id)."""
    r = post(client, "/api/sessions", CSV, "내 BOM.csv")
    sid, jid = r.json()["session_id"], r.json()["job_id"]
    jobs.claim(client.ci)
    jobs.run_match_job(client.ci, jid, lambda ci: FakeEmbedder())
    return client, sid


def patch(client, sid, row, body):
    return client.patch(f"/api/sessions/{sid}/rows/{row}", json=body)


def row_of(client, sid, row):
    return next(r for r in client.get(f"/api/sessions/{sid}").json()["rows"] if r["id"] == row)


def decisions(ci, sid):
    with psycopg.connect(ci) as conn:
        return conn.execute(
            "SELECT d.action, d.outcome, d.from_status, d.to_status, d.reason, d.payload FROM review_decision d JOIN bom_item i USING (item_id) "
            "WHERE i.project_id = %s ORDER BY d.decision_id", (sid,)).fetchall()


@needs_db
def test_default_selection_and_conversion(sess):
    client, sid = sess
    rows = SessionPayload.model_validate(client.get(f"/api/sessions/{sid}").json()).model_dump()["rows"]
    kg, ea, m2, bad, excl = rows
    assert kg["selected"]["manual"] is False and kg["selected"]["score"] is not None and kg["conversion"]["status"] == "ok"
    assert ea["conversion"]["missing"] == ["개당 중량"] and m2["conversion"]["missing"] == ["두께", "밀도"]
    assert bad["conversion"]["missing"] == ["단위"] and excl["review"]["status"] == "excluded"


@needs_db
def test_out_of_range_unknown_and_validation_leave_row_unchanged_but_recorded(sess):
    client, sid = sess
    before = row_of(client, sid, 2)
    assert_error(patch(client, sid, 2, {"thickness_mm": 0}), 400, "OUT_OF_RANGE", {"field": "thickness_mm"})
    assert_error(patch(client, sid, 2, {"thickness_mm": 1000.5}), 400, "OUT_OF_RANGE", {"field": "thickness_mm"})
    assert_error(patch(client, sid, 2, {"density_kg_m3": -1}), 400, "OUT_OF_RANGE", {"field": "density_kg_m3"})
    assert_error(patch(client, sid, 2, {"unit_mass_kg": 1001, "note": "x"}), 400, "OUT_OF_RANGE", {"field": "unit_mass_kg"})
    assert_error(patch(client, sid, 2, {"selected_id": "list-999"}), 400, "UNKNOWN_DATASET")
    assert_error(patch(client, sid, 2, {"review_status": "done"}), 422, "VALIDATION_ERROR")
    assert_error(patch(client, sid, 2, {"review_status": "approved"}), 409, "CONVERSION_INCOMPLETE")
    assert row_of(client, sid, 2) == before  # 아무것도 바뀌지 않았다
    recs = decisions(client.ci, sid)
    assert len(recs) == 7 and {r[1] for r in recs} == {"rejected"} and recs[3][5] == {"unit_mass_kg": 1001, "note": "x"}
    assert_error(patch(client, sid, 99, {"note": "x"}), 404, "ROW_NOT_FOUND")
    assert_error(patch(client, "abcdefabcdef", 1, {"note": "x"}), 404, "SESSION_NOT_FOUND")


@needs_db
def test_direct_selection_is_manual_without_score(sess):
    client, sid = sess
    row = row_of(client, sid, 1)
    top3 = {c["dataset_id"] for c in row["candidates"]}
    outside = next(f"list-{n}" for n in range(5, 27) if f"list-{n}" not in top3)
    r = patch(client, sid, 1, {"selected_id": outside})
    body = RowUpdateResponse.model_validate(r.json()).model_dump()
    sel = body["row"]["selected"]
    assert (sel["dataset_id"], sel["manual"], sel["score"], sel["confidence"]) == (outside, True, None, None)
    assert body["row"]["conversion"]["status"] == "ok"  # 합성 데이터셋 단위 kg
    inside = row["candidates"][1]["dataset_id"]
    sel = patch(client, sid, 1, {"selected_id": inside}).json()["row"]["selected"]
    assert sel["manual"] is False and sel["score"] == row["candidates"][1]["score"]
    cv = RowUpdateResponse.model_validate(patch(client, sid, 1, {"selected_id": "list-4"}).json()).row.conversion  # 단위 없는 합성 데이터셋
    assert (cv.status, cv.unit) == ("incompatible", None)
    r = patch(client, sid, 1, {"clear_selection": True}).json()["row"]
    assert (r["selected_id"], r["selected"], r["conversion"]) == (None, None, None)
    resp = patch(client, sid, 1, {"review_status": "approved"})
    assert_error(resp, 409, "CONVERSION_INCOMPLETE")
    assert "데이터셋을 선택하지 않았습니다" in resp.json()["detail"]


@needs_db
def test_unit_mass_then_approve_then_selection_change_releases(sess):
    client, sid = sess
    resp = patch(client, sid, 2, {"review_status": "approved"})
    assert_error(resp, 409, "CONVERSION_INCOMPLETE")
    assert "개당 중량이 필요합니다" in resp.json()["detail"]
    r = patch(client, sid, 2, {"unit_mass_kg": 0.35}).json()["row"]
    # 응답 모델(Overrides)이 보내지 않은 값도 null 키로 내보낸다
    assert ({k: v for k, v in r["overrides"].items() if v is not None}, r["conversion"]["status"], r["conversion"]["value"],
            r["conversion"]["method"]) == ({"unit_mass_kg": 0.35}, "converted", 7.0, "unit_mass")
    body = patch(client, sid, 2, {"review_status": "approved", "note": "도면 확인"}).json()
    assert body["row"]["review"] == {"status": "approved", "note": "도면 확인"} and body["summary"]["review"]["approved"] == 1
    other = next(c["dataset_id"] for c in body["row"]["candidates"] if c["dataset_id"] != body["row"]["selected_id"])
    r = patch(client, sid, 2, {"selected_id": other}).json()["row"]
    assert r["review"]["status"] == "pending"  # 승인 해제
    last = decisions(client.ci, sid)[-1]
    assert (last[0], last[1], last[2], last[3], last[4]) == ("select", "applied", "approved", "pending", "선택 변경으로 승인 해제")


@needs_db
def test_reject_exclude_and_summary(sess):
    client, sid = sess
    s = patch(client, sid, 3, {"review_status": "rejected", "note": "규격 확인 필요"}).json()["summary"]
    assert s["review"]["rejected"] == 1
    s = patch(client, sid, 4, {"review_status": "excluded"}).json()["summary"]
    assert (s["review"]["excluded"], s["review"]["pending"], s["needs_review"]) == (2, 2, 2)


@needs_db
def test_explain_saved_and_reset(sess):
    client, sid = sess
    e = client.post(f"/api/sessions/{sid}/rows/2/explain")
    assert e.status_code == 200
    ex = Explanation.model_validate(e.json())
    assert ex.source == "rule" and "선택 근거" in ex.text and "확률이 아닙니다" in ex.text and "환산:" in ex.text
    assert ex.questions == ["개당 중량 값을 확인해 주세요"]
    assert row_of(client, sid, 2)["explanation"] == e.json()  # 저장됨
    assert patch(client, sid, 2, {"note": "메모만"}).json()["row"]["explanation"] == e.json()  # 메모만 바꾸면 유지
    assert patch(client, sid, 2, {"unit_mass_kg": 0.5}).json()["row"]["explanation"] is None  # 환산 입력이 바뀌면 초기화
    assert_error(client.post(f"/api/sessions/{sid}/rows/42/explain"), 404, "ROW_NOT_FOUND")


@needs_db
def test_concurrent_patch_waits_for_row_lock(sess):
    client, sid = sess
    holder = psycopg.connect(client.ci)
    holder.execute("SELECT 1 FROM bom_item WHERE project_id = %s AND row_no = 1 FOR UPDATE", (sid,))
    result = {}
    t = threading.Thread(target=lambda: result.setdefault("r", patch(client, sid, 1, {"note": "두 번째"})))
    t.start()
    time.sleep(1.0)
    assert t.is_alive()  # 첫 트랜잭션이 행을 잡고 있는 동안 기다린다
    holder.execute("UPDATE bom_item SET review_note = '첫 번째' WHERE project_id = %s AND row_no = 1", (sid,))
    holder.commit()
    holder.close()
    t.join(timeout=10)
    assert result["r"].status_code == 200 and result["r"].json()["row"]["review"]["note"] == "두 번째"


@needs_db
def test_export_xlsx_csv_and_409(sess):
    client, sid = sess
    assert_error(client.get(f"/api/sessions/{sid}/export"), 409, "NO_APPROVED_ROWS")
    patch(client, sid, 1, {"review_status": "approved"})
    r = client.get(f"/api/sessions/{sid}/export")
    assert r.status_code == 200
    cd = r.headers["content-disposition"]
    assert 'filename="LCI-Match_BOM_' in cd and "filename*=UTF-8''LCI-Match_%EB%82%B4%20BOM_" in cd and cd.endswith(".xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["승인", "미승인·제외", "정보"]
    # X1: 1행 그룹 띠, 2행 열 이름, 3행부터 데이터(+ 합계 한 줄)
    assert (wb["승인"].max_row, wb["미승인·제외"].max_row) == (4, 6)  # 띠·머리 2행 + 승인 1 + 합계 1 / 띠·머리 2행 + 나머지 4
    assert wb["승인"].freeze_panes == "E3" and wb["승인"]["A1"].value == "① 내가 올린 BOM (원본)" and wb["승인"]["A2"].value == "원본 행"
    head = [c.value for c in wb["승인"][2]]
    assert head[:4] == ["원본 행", "부품번호", "구분", "품명"] and "데이터셋 역할" in head
    assert [c.value for c in wb["승인"][3]][1:3] == [None, None]  # 구분 열이 없는 BOM: 빈 칸
    info = {row[0].value: row[1].value for row in wb["정보"].iter_rows(min_row=2)}
    assert info["세션 ID"] == sid and info["안내"] == "이 결과는 사람이 승인한 매칭이며 인증용 계산이 아닙니다."
    c = client.get(f"/api/sessions/{sid}/export", params={"format": "csv"})
    assert c.content.startswith(b"\xef\xbb\xbf") and c.headers["content-disposition"].endswith(".csv")
    lines = list(csv.reader(io.StringIO(c.content.decode("utf-8-sig"))))
    assert lines[0][:4] == ["원본 행", "부품번호", "구분", "품명"] and len(lines) == 2 and lines[1][1:3] == ["", ""]
    assert len(decisions(client.ci, sid)) == 1  # 내보내기는 기록하지 않는다


@needs_db
def test_template_reads_back(client):  # noqa: F811
    r = client.get("/api/template.xlsx")
    assert r.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["BOM", "작성 안내"] and wb["BOM"]["C2"].comment is not None
    header, rows = read_bom("BOM_양식.xlsx", r.content)
    assert header == ["부품번호", "구분", "품명", "규격", "단위", "수량"] and len(rows) == 3
    assert wb["BOM"]["A1"].comment is not None and wb["BOM"]["B1"].comment is not None  # 앞의 두 열은 "선택" 메모
    p = post(client, "/api/sessions/preview", r.content, "양식.xlsx")
    assert p.status_code == 200 and {m["field"]: (m["method"], m["required"]) for m in p.json()["column_mapping"]} == {
        "품명": ("auto", True), "규격": ("auto", False), "단위": ("auto", True), "수량": ("auto", True),
        "구분": ("auto", False), "부품번호": ("auto", False)}
    s = post(client, "/api/sessions", r.content, "양식.xlsx")  # 예시 행을 지우지 않고 올려도 구분이 읽힌다
    assert s.status_code == 202


@needs_db
def test_review_decision_is_append_only(sess):
    client, sid = sess
    patch(client, sid, 1, {"note": "a"})
    with psycopg.connect(client.ci, autocommit=True) as conn:
        with pytest.raises(errors.InsufficientPrivilege):
            conn.execute("UPDATE review_decision SET note = 'b'")
        with pytest.raises(errors.InsufficientPrivilege):
            conn.execute("DELETE FROM review_decision")
        assert conn.execute("SELECT count(*) FROM review_decision").fetchone()[0] == 1
        conn.execute("DELETE FROM bom_project WHERE project_id = %s", (sid,))  # 만료 정리와 같은 경로
        assert conn.execute("SELECT count(*) FROM review_decision").fetchone()[0] == 0


@needs_db
def test_0005_backfills_t9a_rows(tempdb, tmp_path):  # noqa: F811
    old = tmp_path / "m"
    old.mkdir()
    for f in sorted(MIGRATIONS.glob("*.sql"))[:4]:  # 0001~0004 (T9a 시점)
        shutil.copy(f, old)
    migrate(tempdb, old)
    with psycopg.connect(tempdb) as conn:
        conn.execute("INSERT INTO bom_project (project_id, source_file, expires_at) VALUES ('abcdefABCDEF', 'a.csv', now() + interval '1 day')")
        for n, status in ((1, "대상"), (2, "제외")):
            conn.execute("INSERT INTO bom_item (project_id, row_no, source_row, raw_name, raw_spec, raw_unit, raw_qty, clean_name, clean_spec, "
                         "clean_unit_display, norm_text, query_key, target_status) VALUES ('abcdefABCDEF', %s, %s, 'a', '', 'kg', '1', 'a', '', "
                         "'kg', 'a', 'k', %s)", (n, n + 1, status))
    migrate(tempdb)  # 0005 적용
    with psycopg.connect(tempdb) as conn:
        assert conn.execute("SELECT row_no, review_status, selected_id, selected_manual, review_note FROM bom_item ORDER BY row_no").fetchall() == [
            (1, "pending", None, None, ""), (2, "excluded", None, None, "")]
