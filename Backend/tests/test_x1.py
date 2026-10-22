"""X1 배출계수: XML 조각에서 추출, 적재 정규화·이슈, A02·A06 carbon, 내보내기(xlsx 서식·추정 배출량·합계·csv). 임시 DB + 가짜 임베더."""
import csv
import io
from decimal import Decimal

import openpyxl
import psycopg
import pytest

from app.bom import review
from app.carbon import carbon_of
from app.ingest import build, gwp_values
from app.schemas import DatasetDetail, SessionPayload
from test_ingest import KEITI, LIST_COLS, MAP_COLS  # noqa: F401
from test_migrate import needs_db, tempdb  # noqa: F401
from test_review_api import patch, sess  # noqa: F401  (fixture 재사용)
from test_bom_match import bomdb  # noqa: F401
from test_sessions_api import client  # noqa: F401
from tools.link_files import read_lcia


def lcia(label, amount, version="04.00.016"):
    return (f'<LCIAResult><referenceToLCIAMethodDataSet type="LCIA method data set" refObjectId="u-{label}" version="{version}">'
            f'<common:shortDescription xml:lang="en">{label}</common:shortDescription></referenceToLCIAMethodDataSet>'
            f"<meanAmount>{amount}</meanAmount></LCIAResult>")


FULL = "<LCIAResults>" + "".join([
    lcia("Acidification", "0.0102", "01.03.009"), lcia("Climate change", "0.98384"), lcia("Climate change-Biogenic", "0.02791", "01.00.005"),
    lcia("Climate change-Fossil", "0.95337", "01.00.008"), lcia("Climate change-Land use and land use change", "0.00255", "01.00.005"),
    lcia("Climate change-Other", "9"), lcia("Ecotoxicity, freshwater", "9.05")]) + "</LCIAResults>"


# ---------- ① 추출 ----------
def test_read_lcia_reads_only_the_four_climate_items_and_method_version():
    got = read_lcia(FULL)
    assert got == {"total": "0.98384", "method_version": "04.00.016", "biogenic": "0.02791", "fossil": "0.95337", "luluc": "0.00255"}


def test_read_lcia_none_when_no_results_or_no_climate_item():
    assert read_lcia("<exchanges></exchanges>") is None
    assert read_lcia("<LCIAResults>" + lcia("Acidification", "1") + "</LCIAResults>") is None


def test_read_lcia_missing_amount_is_kept_as_none_not_zero():
    text = "<LCIAResults>" + lcia("Climate change", "1") + (lcia("Climate change-Fossil", "1").replace("<meanAmount>1</meanAmount>", "")) + "</LCIAResults>"
    assert read_lcia(text)["fossil"] is None


def test_gwp_values_normal_reference_flow_1():
    values, issues = gwp_values(read_lcia(FULL), Decimal("1"))
    assert issues == []
    assert values == {"gwp_total": Decimal("0.98384"), "gwp_fossil": Decimal("0.95337"), "gwp_biogenic": Decimal("0.02791"),
                      "gwp_luluc": Decimal("0.00255"), "gwp_method": "Climate change (04.00.016)"}


def test_gwp_values_normalizes_per_reference_unit_when_amount_is_1000_kg():
    values, issues = gwp_values(read_lcia(FULL), Decimal("1000"))
    assert issues == []
    assert values["gwp_total"] == Decimal("0.00098384") and values["gwp_fossil"] == Decimal("0.00095337")


@pytest.mark.parametrize("lcia_value", [None, {}, {"total": None}])
def test_gwp_values_no_result_is_null_without_issue(lcia_value):
    values, issues = gwp_values(lcia_value, Decimal("1"))
    assert issues == [] and set(values.values()) == {None}


@pytest.mark.parametrize("ref", [None, Decimal("0")])
def test_gwp_values_no_reference_amount_is_null_with_issue(ref):
    values, issues = gwp_values(read_lcia(FULL), ref)
    assert set(values.values()) == {None} and [i[0] for i in issues] == ["gwp_no_ref"]


def test_gwp_values_sum_mismatch_is_an_issue_but_keeps_values():
    text = FULL.replace("<meanAmount>0.95337</meanAmount>", "<meanAmount>0.5</meanAmount>")  # 세 항목 합이 0.53 → 합계 0.98 과 1% 넘게 다름
    values, issues = gwp_values(read_lcia(text), Decimal("1"))
    assert [(i[0], i[1]) for i in issues] == [("gwp_sum_mismatch", "info")]
    assert values["gwp_total"] == Decimal("0.98384") and values["gwp_fossil"] == Decimal("0.5")


def test_gwp_values_sum_within_one_percent_is_fine():
    text = FULL.replace("<meanAmount>0.95337</meanAmount>", "<meanAmount>0.95000</meanAmount>")  # 합 0.98046 vs 0.98384: 0.34%
    assert gwp_values(read_lcia(text), Decimal("1"))[1] == []


def test_build_puts_normalized_values_and_issues_into_dataset_rows():
    lst = [{"lci_database_sn": str(n), "category": "금속", "name_ko": f"강판{n}", "name_en": f"Steel {n}", "org": KEITI, "year": "2022", "version": "1.0.0",
            "zip_file": f"LCI-2022-00{n}.zip", "excel_file": ""} for n in (1, 2, 3)]
    mp = [{"file": f"LCI-2022-00{n}.zip", "sha": str(n) * 64, "kind": "xml", "list_row": str(n), "status": "exact", "top3": "", "name_check": "0.9"} for n in (1, 2, 3)]

    def feat(n, amount, lc):
        return {"sha": str(n) * 64, "kind": "xml", "uuid": f"{n}" * 8, "reference_year": 2022, "geography": "KR", "dataset_type": "LCI result",
                "allocation": "a", "cutoff": "c", "tech": "t", "lcia": lc, "ref_flow": {"name": "f", "amount": amount, "unit": "kg"}}
    features = {"LCI-2022-001.zip": feat(1, "1,000", read_lcia(FULL)), "LCI-2022-002.zip": feat(2, "1.0", None),
                "LCI-2022-003.zip": feat(3, None, read_lcia(FULL))}
    datasets, issues = build(lst, mp, features, seed={})
    by = {d["dataset_id"]: d for d in datasets}
    assert by["list-1"]["gwp_total"] == Decimal("0.00098384")  # 1000 kg 당 → 1 kg 당
    assert by["list-2"]["gwp_total"] is None and by["list-2"]["gwp_method"] is None  # 영향평가 결과 없음 → NULL
    assert by["list-3"]["gwp_total"] is None  # 기준 흐름 양 없음 → NULL + 이슈
    assert [(i[0], i[1]) for i in issues if i[1].startswith("gwp")] == [("list-3", "gwp_no_ref")]


# ---------- ② A02·A06 carbon ----------
def set_gwp(ci, rows):
    with psycopg.connect(ci) as conn:
        for dataset_id, total, fossil, bio, luluc in rows:
            conn.execute("UPDATE lci_dataset SET gwp_total = %s, gwp_fossil = %s, gwp_biogenic = %s, gwp_luluc = %s, gwp_method = 'Climate change (04.00.016)' "
                         "WHERE dataset_id = %s", (total, fossil, bio, luluc, dataset_id))


def test_carbon_of_none_without_total_and_float_values():
    assert carbon_of(None, 1, 1, 1, "m", "kg") is None
    c = carbon_of(Decimal("2.0"), Decimal("1.5"), None, Decimal("0.1"), None, None)
    assert c == {"total": 2.0, "fossil": 1.5, "biogenic": None, "land_use": 0.1, "unit": "kg CO₂e", "per": None, "method": "Climate change"}


@needs_db
def test_a06_and_a02_carry_carbon_or_null(sess):  # noqa: F811
    client, sid = sess
    set_gwp(client.ci, [("list-5", Decimal("2.0"), Decimal("1.5"), Decimal("0.4"), Decimal("0.1"))])
    d = client.get("/api/datasets/list-5").json()
    DatasetDetail.model_validate(d)
    assert d["carbon"] == {"total": 2.0, "fossil": 1.5, "biogenic": 0.4, "land_use": 0.1, "unit": "kg CO₂e", "per": "kg", "method": "Climate change (04.00.016)"}
    assert client.get("/api/datasets/list-6").json()["carbon"] is None  # 값이 없으면 null (키는 있다)
    assert "carbon" not in client.get("/api/datasets", params={"q": "데이터셋5"}).json()["items"][0]  # A05 목록은 제외
    payload = SessionPayload.model_validate(client.get(f"/api/sessions/{sid}").json()).model_dump()
    cands = [c for r in payload["rows"] for c in r["candidates"]]
    assert cands and all((c["carbon"] is not None) == (c["dataset_id"] == "list-5") for c in cands)


# ---------- ③ 내보내기 ----------
def approve_three(client, sid):
    """1: 냉연강판 2.4 kg → list-5(계수 있음) / 2: 볼트 20 EA × 0.35 kg = 7.0 kg → list-6(계수 없음) / 3: 알루미늄 판 0.8 m² × 2 mm × 2700 = 4.32 kg → list-7(계수 있음)."""
    set_gwp(client.ci, [("list-5", Decimal("2.0"), Decimal("1.5"), Decimal("0.4"), Decimal("0.1")),
                        ("list-7", Decimal("0.5"), Decimal("0.3"), Decimal("0.15"), Decimal("0.05"))])
    assert patch(client, sid, 1, {"selected_id": "list-5", "review_status": "approved"}).status_code == 200
    assert patch(client, sid, 2, {"selected_id": "list-6", "unit_mass_kg": 0.35, "review_status": "approved"}).status_code == 200
    assert patch(client, sid, 3, {"selected_id": "list-7", "thickness_mm": 2, "density_kg_m3": 2700, "review_status": "approved"}).status_code == 200


def thick_left(cell):
    return cell.border.left.style == "thick"


@needs_db
def test_export_xlsx_groups_colors_borders_freeze_filter_estimate_and_subtotal(sess):  # noqa: F811
    client, sid = sess
    approve_three(client, sid)
    wb = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/sessions/{sid}/export").content))
    assert wb.sheetnames == ["승인", "미승인·제외", "정보"]
    ws = wb["승인"]
    # 그룹 띠: 1행 셀 병합(① A:G, ② H:O, ③ P:R, ④ S:X, ⑤ Y:Z), 글자는 표기하고 굵게
    assert sorted(str(r) for r in ws.merged_cells.ranges if r.min_row == 1) == ["A1:G1", "H1:O1", "P1:R1", "S1:X1", "Y1:Z1"]
    assert [ws[c].value for c in ("A1", "H1", "P1", "S1", "Y1")] == [
        "① 내가 올린 BOM (원본)", "② 연결한 국가 LCI DB 데이터셋", "③ 환산", "④ 배출계수·추정 배출량 (참고용)", "⑤ 검토"]
    assert all(ws[c].font.bold for c in ("A1", "H1", "P1", "S1", "Y1", "A2", "H2", "S2"))
    assert [c.value for c in ws[2]][:8] == ["원본 행", "부품번호", "구분", "품명", "규격", "수량", "단위", "데이터셋 ID"]
    # 머리 색·본문 색(그룹별)
    colors = {"A": ("BBD6F5", "EEF5FD"), "H": ("BFE5C8", "F0F9F2"), "P": ("F6E0A8", "FDF7E6"), "S": ("D9C8F0", "F6F1FC"), "Y": ("DDDDDD", "FFFFFF")}
    for col, (head, body) in colors.items():
        assert ws[f"{col}1"].fill.start_color.rgb[-6:] == head and ws[f"{col}2"].fill.start_color.rgb[-6:] == head, col
        assert ws[f"{col}3"].fill.start_color.rgb[-6:] == body, col
    # 그룹 사이 굵은 세로 테두리: 각 그룹 첫 열의 왼쪽, 마지막 열의 오른쪽
    assert all(thick_left(ws[f"{c}{r}"]) for c in "AHPSY" for r in (1, 2, 3, 5))
    assert ws["Z3"].border.right.style == "thick" and ws["G3"].border.right.style == "thick"
    assert ws.freeze_panes == "E3" and ws.auto_filter.ref == "A2:Z5"
    # 추정 배출량 = 환산 값 × 총 배출계수 (손 계산: 2.4×2.0 = 4.8 / 7.0×(계수 없음) / 4.32×0.5 = 2.16)
    head = [c.value for c in ws[2]]
    coef, est, state = head.index("배출계수(kg CO₂e/기준 단위)") + 1, head.index("추정 배출량(kg CO₂e)") + 1, head.index("계수 상태") + 1
    assert ws.cell(3, est).value == pytest.approx(4.8) and ws.cell(5, est).value == pytest.approx(2.16)
    assert ws.cell(4, coef).value is None and ws.cell(4, est).value is None  # 계수 없는 행은 빈 칸(0 금지)
    assert [ws.cell(r, state).value for r in (3, 4, 5)] == ["계산됨", "계수 없음(이 데이터셋에 영향평가 결과가 없음)", "계산됨"]
    assert [ws.cell(3, c).value for c in range(coef, coef + 4)] == [2.0, 1.5, 0.4, 0.1]
    assert ws.cell(3, coef).number_format == "0.0000" and ws.cell(3, est).number_format == "0.000"
    # 합계(참고용): SUBTOTAL(109) 수식, 필터에 맞는 보이는 행만. 옆 칸에 빠진 행 수
    est_col = ws.cell(2, est).column_letter
    assert ws.cell(6, 1).value == "합계 (참고용)"
    assert ws.cell(6, est).value == f"=SUBTOTAL(109,{est_col}3:{est_col}5)"
    assert ws.cell(6, state).value == "계수가 없는 1행·환산 안 된 0행은 합계에서 빠졌습니다"
    # 미승인·제외 시트: ① + ⑤ 상태
    other = wb["미승인·제외"]
    assert sorted(str(r) for r in other.merged_cells.ranges if r.min_row == 1) == ["A1:G1", "H1:I1"]
    assert (other["A1"].value, other["H1"].value, other["H2"].value) == ("① 내가 올린 BOM (원본)", "⑤ 상태", "상태")
    assert other["A2"].fill.start_color.rgb[-6:] == "BBD6F5" and other["H2"].fill.start_color.rgb[-6:] == "DDDDDD"
    assert other.freeze_panes == "E3" and other.auto_filter.ref.startswith("A2:I")
    # 정보 시트: 기존 행은 그대로, 행 추가
    info = {row[0].value: row[1].value for row in wb["정보"].iter_rows(min_row=2)}
    assert info["세션 ID"] == sid and info["안내"] == review.DISCLAIMER
    assert info["표 구분 안내"] == "① 파랑 = 올린 값, ② 초록 = 국가 DB, ③ 노랑 = 환산, ④ 보라 = 배출계수·추정, ⑤ 회색 = 검토"
    assert info["배출계수 출처"] == "국가 LCI DB 데이터셋의 영향평가 결과, Climate change (04.00.016), 기준 단위 1당"
    assert "참고용 추정이며 인증·검증용 계산이 아닙니다" in info["주의"] and "합계는 전체 배출량이 아닙니다" in info["주의"]
    assert info["계수가 있는 데이터셋 비율"].startswith("2 / 26건(8%)")


@needs_db
def test_export_csv_same_column_order_with_carbon_columns(sess):  # noqa: F811
    client, sid = sess
    approve_three(client, sid)
    c = client.get(f"/api/sessions/{sid}/export", params={"format": "csv"})
    lines = list(csv.reader(io.StringIO(c.content.decode("utf-8-sig"))))
    assert lines[0] == review.APPROVED_COLS and len(lines) == 4 and all(len(r) == 26 for r in lines)
    i = lines[0].index
    assert lines[0][i("배출계수(kg CO₂e/기준 단위)"):i("계수 상태") + 1] == ["배출계수(kg CO₂e/기준 단위)", "화석", "생물기원", "토지이용", "추정 배출량(kg CO₂e)", "계수 상태"]
    assert float(lines[1][i("추정 배출량(kg CO₂e)")]) == pytest.approx(4.8) and lines[2][i("배출계수(kg CO₂e/기준 단위)")] == "" and lines[2][i("추정 배출량(kg CO₂e)")] == ""


def test_estimate_only_when_converted_and_same_unit_and_coefficient():
    ok = {"status": "converted", "value": 7.0, "unit": "kg"}
    assert review.estimate(ok, "kg", 2.0) == (14.0, review.STATE_OK)
    assert review.estimate(ok, "kg", None) == (None, review.STATE_NONE)  # 계수 없음이 먼저
    assert review.estimate({"status": "needs_info", "value": None, "unit": "kg"}, "kg", 2.0) == (None, review.STATE_UNCONVERTED)
    assert review.estimate({"status": "ok", "value": 7.0, "unit": "MJ"}, "kg", 2.0) == (None, review.STATE_UNCONVERTED)  # 환산 단위 ≠ 기준 단위
    assert review.estimate({"status": "converted", "value": 4.0, "unit": "EA"}, "Item(s)", 0.5) == (2.0, review.STATE_OK)  # 개수 단위 표기 Item(s) = EA
    assert review.estimate(None, "kg", 2.0) == (None, review.STATE_UNCONVERTED)
